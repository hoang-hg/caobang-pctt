"""Đọc tệp nhập (CSV / Excel .xlsx / GeoJSON) thành các dòng thô + chuẩn hoá giá trị. Hàm thuần — không cần CSDL."""

from __future__ import annotations

import csv
import io
import json
import math
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from app.services.data_import.specs import MAX_FILE_BYTES, MAX_ROWS


class ImportFileError(ValueError):
    """Tệp không đọc được (sai định dạng, quá lớn…) — lỗi của cả tệp, không phải từng dòng."""


@dataclass
class RawRow:
    number: int  # số dòng hiển thị cho người dùng (dòng tiêu đề CSV/Excel = 1; GeoJSON: thứ tự feature từ 1)
    values: dict[str, Any]  # khoá đã chuẩn hoá (norm_key)
    geometry: dict | None = None  # GeoJSON geometry (chỉ tệp GeoJSON)
    extra_columns: list[str] = field(default_factory=list)


def strip_accents(text: str) -> str:
    text = unicodedata.normalize("NFD", text.replace("đ", "d").replace("Đ", "D"))
    return "".join(c for c in text if unicodedata.category(c) != "Mn")


def norm_key(text: str) -> str:
    """'Vĩ độ' → 'vi_do', 'Sức chứa (người)' → 'suc_chua_nguoi', 'MA XA' → 'ma_xa'."""
    return re.sub(r"[^a-z0-9]+", "_", strip_accents(str(text)).lower()).strip("_")


def _cell(value: Any) -> Any:
    """Giá trị ô Excel → chuỗi (giữ date/datetime); 1200.0 → '1200'."""
    if value is None or isinstance(value, datetime | date):
        return value
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _rows_from_table(header: list[Any], body: list[tuple[int, list[Any]]]) -> list[RawRow]:
    keys = [norm_key(h) if h not in (None, "") else "" for h in header]
    if not any(keys):
        raise ImportFileError("Không tìm thấy dòng tiêu đề (dòng 1 phải là tên cột)")
    rows = []
    for number, cells in body:
        values = {k: _cell(v) for k, v in zip(keys, cells, strict=False) if k}
        if all(v in (None, "") for v in values.values()):
            continue  # bỏ dòng trống
        rows.append(RawRow(number, values))
    return rows


def read_csv(data: bytes) -> list[RawRow]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ImportFileError(
            "Tệp CSV không phải mã UTF-8 — trong Excel chọn Lưu thành → “CSV UTF-8 (phân cách bằng dấu phẩy)”"
        ) from exc
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    lines = list(reader)
    if not lines:
        raise ImportFileError("Tệp trống")
    return _rows_from_table(lines[0], [(i + 2, line) for i, line in enumerate(lines[1:])])


def read_xlsx(data: bytes) -> list[RawRow]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # zip hỏng, không phải xlsx…
        raise ImportFileError("Không đọc được tệp Excel (.xlsx)") from exc
    ws = wb.worksheets[0]
    # Dừng khi quá giới hạn: .xlsx là zip — vài MB có thể giải nén ra hàng triệu dòng (kể cả dòng trống do định dạng kéo
    # tới cuối trang tính); đọc hết vào bộ nhớ rồi mới đếm thì tiến trình API cạn RAM. Dòng dữ liệu đứng trước dòng trống.
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(row)
        if len(rows) > MAX_ROWS + 1:
            break
    wb.close()
    if not rows:
        raise ImportFileError("Trang tính đầu tiên trống")
    return _rows_from_table(list(rows[0]), [(i + 2, list(r)) for i, r in enumerate(rows[1:])])


def read_geojson(data: bytes) -> list[RawRow]:
    try:
        doc = json.loads(data.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ImportFileError("Tệp GeoJSON không hợp lệ (không đọc được JSON)") from exc
    crs = (
        (((doc.get("crs") or {}).get("properties") or {}).get("name") or "") if isinstance(doc, dict) else ""
    )
    if crs and not re.search(r"(4326|CRS84)$", crs):
        raise ImportFileError(
            f"Hệ toạ độ {crs} chưa hỗ trợ — xuất lại tệp theo WGS84 (EPSG:4326), không dùng VN-2000"
        )
    if isinstance(doc, dict) and doc.get("type") == "FeatureCollection":
        features = doc.get("features") or []
    elif isinstance(doc, dict) and doc.get("type") == "Feature":
        features = [doc]
    else:
        raise ImportFileError("GeoJSON phải là FeatureCollection hoặc Feature")
    rows = []
    for i, feat in enumerate(features, start=1):
        props = (feat or {}).get("properties") or {}
        rows.append(RawRow(i, {norm_key(k): v for k, v in props.items()}, (feat or {}).get("geometry")))
    return rows


def read_file(filename: str, data: bytes) -> list[RawRow]:
    if len(data) > MAX_FILE_BYTES:
        raise ImportFileError(f"Tệp quá lớn (tối đa {MAX_FILE_BYTES // (1024 * 1024)} MB)")
    name = filename.lower()
    if name.endswith(".csv"):
        rows = read_csv(data)
    elif name.endswith(".xlsx"):
        rows = read_xlsx(data)
    elif name.endswith((".geojson", ".json")):
        rows = read_geojson(data)
    else:
        raise ImportFileError("Chỉ nhận tệp .csv, .xlsx, .geojson")
    if not rows:
        raise ImportFileError("Tệp không có dòng dữ liệu nào")
    if len(rows) > MAX_ROWS:
        raise ImportFileError(f"Quá {MAX_ROWS} dòng — chia nhỏ tệp")
    return rows


# ------------------------------------------------------------------ chuẩn hoá giá trị (ném ValueError với lời nhắn)

_INT_RE = re.compile(r"-?\d{1,3}(?:[.,\s]\d{3})+|-?\d+")
_PHONE_RE = re.compile(r"^[0-9 +().-]{8,20}$")
_TRUE = {"co", "c", "x", "true", "1", "yes", "y", "dung", "dang_hoat_dong"}
_FALSE = {"khong", "k", "false", "0", "no", "n", "sai", ""}


def to_int(value: Any) -> int:
    s = str(value).strip()
    if not _INT_RE.fullmatch(s):
        raise ValueError(f"“{s}” không phải số nguyên")
    return int(re.sub(r"[.,\s]", "", s))


def to_float(value: Any) -> float:
    if isinstance(value, int | float) and math.isfinite(value):
        return float(value)
    s = str(value).strip().replace(" ", "")
    if "," in s and "." in s:
        raise ValueError(
            f"“{s}” có cả dấu chấm và dấu phẩy — dùng dấu chấm thập phân, không dùng dấu ngăn nghìn"
        )
    try:
        value = float(s.replace(",", "."))
    except ValueError as exc:
        raise ValueError(f"“{s}” không phải số") from exc
    if not math.isfinite(value):  # float() nhận "nan", "inf" — NaN làm mọi phép so ngưỡng báo động sai
        raise ValueError(f"“{s}” không phải số")
    return value


def to_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"“{s}” không phải ngày (dùng YYYY-MM-DD hoặc DD/MM/YYYY)")


def to_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    s = norm_key(str(value))
    if s in _TRUE:
        return True
    if s in _FALSE:
        return False
    raise ValueError(f"“{value}” không phải có/không")


def to_enum(value: Any, choices: tuple[str, ...]) -> str:
    s = norm_key(str(value))
    if s not in choices:
        raise ValueError(f"“{value}” không hợp lệ — chọn một trong: {', '.join(choices)}")
    return s


def to_phone(value: Any) -> str:
    s = str(value).strip()
    if not _PHONE_RE.match(s):
        raise ValueError(f"“{s}” không phải số điện thoại")
    return s


def to_code(value: Any) -> str:
    s = str(value).strip()
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.\-/]{0,63}", s):
        raise ValueError(f"“{s}” không hợp lệ — mã chỉ gồm chữ không dấu, số, - _ . / (tối đa 64 ký tự)")
    return s


def to_list(value: Any) -> list[str]:
    return [norm_key(p) for p in re.split(r"[;,|]", str(value)) if p.strip()]


def to_code_list(value: Any) -> list[str]:
    """Danh sách mã hoặc tên (xã trong nhóm lọc nhanh): tách theo ; , | hoặc xuống dòng, giữ nguyên chữ (mã viết hoa,
    tên có dấu — không chuẩn hoá như ``to_list``), bỏ mục trùng, giữ thứ tự."""
    return list(dict.fromkeys(p.strip() for p in re.split(r"[;,|\n]", str(value)) if p.strip()))


CONVERTERS = {
    "int": to_int,
    "float": to_float,
    "date": to_date,
    "bool": to_bool,
    "phone": to_phone,
    "code": to_code,
    "list": to_list,
    "code_list": to_code_list,
}
